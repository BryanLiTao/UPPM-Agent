import os
import re
from pathlib import Path

import pandas as pd
import streamlit as st
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity

try:
    from pypdf import PdfReader
except Exception:
    PdfReader = None

try:
    from pptx import Presentation
except Exception:
    Presentation = None

try:
    from docx import Document
except Exception:
    Document = None

try:
    from openpyxl import load_workbook
except Exception:
    load_workbook = None


# ============================================================
# APP CONFIGURATION
# ============================================================

APP_TITLE = "UPP Mathematics Academic Support Agent"

BASE_DIR = Path(__file__).parent
KNOWLEDGE_DIR = BASE_DIR / "knowledge"
SCHEDULE_FILE = BASE_DIR / "data" / "schedule.csv"


# ============================================================
# STREAMLIT CLOUD / LOCAL SETTINGS
# ============================================================

def get_setting(name, default=""):
    """
    Read a setting from:
    1. Streamlit Cloud secrets
    2. Local environment variables
    3. Default value

    This lets the same app.py work both locally and on
    Streamlit Community Cloud.
    """
    try:
        value = st.secrets.get(name, None)
        if value is not None:
            return str(value).strip()
    except Exception:
        pass

    return os.getenv(name, default).strip()


# ============================================================
# TEXT / FILE HELPERS
# ============================================================

def clean_text(value):
    return re.sub(r"\s+", " ", str(value or "")).strip()


def pdf_units(path: Path):
    if PdfReader is None:
        return []

    reader = PdfReader(str(path))
    units = []

    for page_no, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""

        if text.strip():
            units.append(
                {
                    "source": path.name,
                    "location": f"page {page_no}",
                    "text": text,
                }
            )

    return units


def pptx_units(path: Path):
    if Presentation is None:
        return []

    prs = Presentation(str(path))
    units = []

    for slide_no, slide in enumerate(prs.slides, start=1):
        parts = []

        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                parts.append(shape.text)

            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    parts.append(
                        " | ".join(clean_text(cell.text) for cell in row.cells)
                    )

        text = "\n".join(parts)

        if text.strip():
            units.append(
                {
                    "source": path.name,
                    "location": f"slide {slide_no}",
                    "text": text,
                }
            )

    return units


def docx_units(path: Path):
    if Document is None:
        return []

    doc = Document(str(path))
    text = "\n".join(paragraph.text for paragraph in doc.paragraphs)

    if not text.strip():
        return []

    return [
        {
            "source": path.name,
            "location": "document",
            "text": text,
        }
    ]


def xlsx_units(path: Path):
    if load_workbook is None:
        return []

    wb = load_workbook(path, read_only=True, data_only=True)
    units = []

    for ws in wb.worksheets:
        lines = []

        for row in ws.iter_rows(values_only=True):
            values = [clean_text(v) for v in row if v is not None]

            if values:
                lines.append(" | ".join(values))

        if lines:
            units.append(
                {
                    "source": path.name,
                    "location": f"sheet {ws.title}",
                    "text": "\n".join(lines),
                }
            )

    wb.close()
    return units


def file_units(path: Path):
    suffix = path.suffix.lower()

    if suffix == ".pdf":
        return pdf_units(path)

    if suffix == ".pptx":
        return pptx_units(path)

    if suffix == ".docx":
        return docx_units(path)

    if suffix in {".xlsx", ".xlsm"}:
        return xlsx_units(path)

    if suffix == ".csv":
        df = pd.read_csv(path)
        return [
            {
                "source": path.name,
                "location": "CSV",
                "text": df.astype(str).to_csv(index=False),
            }
        ]

    if suffix in {".txt", ".md", ".tex"}:
        return [
            {
                "source": path.name,
                "location": "document",
                "text": path.read_text(
                    encoding="utf-8",
                    errors="ignore",
                ),
            }
        ]

    return []


def chunk_units(units, chunk_size=1400, overlap=180):
    chunks = []

    for unit in units:
        text = clean_text(unit["text"])
        start = 0

        while start < len(text):
            end = min(len(text), start + chunk_size)
            piece = text[start:end].strip()

            if piece:
                chunks.append(
                    {
                        "source": unit["source"],
                        "location": unit["location"],
                        "text": piece,
                    }
                )

            if end >= len(text):
                break

            start = max(start + 1, end - overlap)

    return chunks


# ============================================================
# KNOWLEDGE BASE
# ============================================================

@st.cache_data(show_spinner=False)
def build_knowledge_base(folder_string):
    folder = Path(folder_string)

    chunks = []
    files = []
    errors = []

    supported = {
        ".pdf",
        ".pptx",
        ".docx",
        ".xlsx",
        ".xlsm",
        ".csv",
        ".txt",
        ".md",
        ".tex",
    }

    if not folder.exists():
        return [], [], [f"Folder not found: {folder}"]

    for path in sorted(folder.rglob("*")):
        if not path.is_file():
            continue

        if path.suffix.lower() not in supported:
            continue

        try:
            units = file_units(path)
            chunks.extend(chunk_units(units))
            files.append(path.name)

        except Exception as exc:
            errors.append(f"{path.name}: {exc}")

    return chunks, files, errors


def retrieve(query, chunks, top_k=6):
    if not chunks:
        return []

    corpus = [chunk["text"] for chunk in chunks]

    try:
        vectorizer = TfidfVectorizer(
            stop_words="english",
            ngram_range=(1, 2),
        )

        matrix = vectorizer.fit_transform(corpus + [query])
        scores = cosine_similarity(
            matrix[-1],
            matrix[:-1],
        ).flatten()

    except ValueError:
        return []

    order = scores.argsort()[::-1][:top_k]

    return [
        {
            **chunks[i],
            "score": float(scores[i]),
        }
        for i in order
        if scores[i] > 0
    ]


# ============================================================
# AI ANSWER
# ============================================================

def call_ai(question, retrieved):
    """
    OpenAI-compatible Chat Completions call.

    Streamlit Cloud Secrets:
        API123_API_KEY
        API123_BASE_URL
        API123_MODEL
    """

    api_key = get_setting("API123_API_KEY")

    if not api_key:
        return None

    base_url = get_setting(
        "API123_BASE_URL",
        "https://api123.top/v1",
    )

    default_model = get_setting(
        "API123_MODEL",
        "gpt-5.5",
    )

    model = st.session_state.get(
        "api123_model",
        default_model,
    )

    try:
        from openai import OpenAI

        client = OpenAI(
            api_key=api_key,
            base_url=base_url,
        )

        context = "\n\n".join(
            f"[{item['source']} — {item['location']}]\n{item['text']}"
            for item in retrieved
        )

        prompt = f"""
You are a support agent for UPP Mathematics students.

IMPORTANT SOURCE RULES
- Use ONLY the supplied approved course extracts for factual course information.
- Never invent dates, assessment rules, venues, trainer details, deadlines, or assessment scope.
- If the supplied sources do not support an answer, say so clearly.
- For schedule and assessment questions, prefer timetable and briefing-slide evidence.

MATHEMATICS RULES
- Explain clearly at UPP Mathematics level.
- Render inline mathematics using $...$.
- Render displayed mathematics using $$...$$.
- Use proper LaTeX notation.
- Do NOT repeat garbled PDF-extraction symbols.
- Reconstruct a formula only when the supplied context supports it.
- If PDF extraction is unclear, explain the concept using clean notation without pretending unreadable symbols were visible.

STUDENT PRIVACY RULE
- This is a student-facing application.
- Do not mention internal staff data, risk labels, support flags, or student monitoring information.

APPROVED SOURCE EXTRACTS
{context}

STUDENT QUESTION
{question}

Give a concise, student-friendly answer.

For mathematics questions, explain the working clearly and include a short worked example when useful.

For course-information questions, answer only from the approved extracts.

At the end, include a short "Source" line giving the filename and page/slide used.
""".strip()

        response = client.chat.completions.create(
            model=model,
            messages=[
                {
                    "role": "system",
                    "content": (
                        "You are a careful UPP Mathematics learning-support assistant. "
                        "Course facts must remain grounded in the supplied approved sources."
                    ),
                },
                {
                    "role": "user",
                    "content": prompt,
                },
            ],
            temperature=0.2,
        )

        content = response.choices[0].message.content

        if content:
            return content.strip()

        return "The AI returned an empty response."

    except Exception as exc:
        return f"AI call unavailable: {exc}"


# ============================================================
# RETRIEVAL-ONLY FALLBACK
# ============================================================

def retrieval_answer(question, retrieved):
    if not retrieved:
        return (
            "I could not find enough information about this in the approved "
            "UPP Mathematics course materials. Please check with your trainer."
        )

    sections = []

    for item in retrieved[:4]:
        extract = re.sub(
            r"[\x00-\x08\x0b\x0c\x0e-\x1f]",
            "",
            item["text"],
        )

        short_extract = extract[:850]

        if len(extract) > 850:
            short_extract += "…"

        sections.append(
            f"**{item['source']} — {item['location']}**\n\n"
            f"{short_extract}"
        )

    return (
        "I found the following relevant information in the approved "
        "UPP Mathematics course materials.\n\n"
        + "\n\n---\n\n".join(sections)
    )


# ============================================================
# PAGE
# ============================================================

st.set_page_config(
    page_title=APP_TITLE,
    page_icon="🎓",
    layout="wide",
)

st.title("🎓 UPP Mathematics Support Agent")

st.caption(
    "Ask questions about UPP Mathematics concepts, programme dates, "
    "assessments and preparation."
)


# ============================================================
# SIDEBAR
# ============================================================

with st.sidebar:
    st.header("UPP Mathematics")

    if st.button(
        "🔄 Refresh course materials",
        use_container_width=True,
    ):
        build_knowledge_base.clear()
        st.rerun()

    st.divider()

    default_model = get_setting(
        "API123_MODEL",
        "gpt-5.5",
    )

    st.text_input(
        "AI model",
        value=default_model,
        key="api123_model",
        help="Use a model name available from your configured AI provider.",
    )

    if get_setting("API123_API_KEY"):
        st.success("AI mode enabled")
    else:
        st.info(
            "AI key not detected. The app will use retrieval-only mode."
        )


# ============================================================
# LOAD COURSE MATERIALS
# ============================================================

chunks, files_read, errors = build_knowledge_base(
    str(KNOWLEDGE_DIR)
)

if not KNOWLEDGE_DIR.exists():
    st.error(
        "The knowledge folder was not found. "
        "Please add a folder named 'knowledge' beside app.py."
    )

elif not files_read:
    st.warning(
        "No supported course-material files were found in the knowledge folder."
    )

if errors:
    with st.expander("Course-material loading notes"):
        for error in errors:
            st.warning(error)


# ============================================================
# LOAD SCHEDULE
# ============================================================

schedule = None

if SCHEDULE_FILE.exists():
    try:
        schedule = pd.read_csv(SCHEDULE_FILE)
        schedule["Date"] = pd.to_datetime(
            schedule["Date"],
            errors="coerce",
        ).dt.date

    except Exception as exc:
        st.warning(
            f"The programme schedule could not be loaded: {exc}"
        )
else:
    st.warning(
        "The schedule file was not found at data/schedule.csv."
    )


# ============================================================
# STUDENT AGENT
# ============================================================

st.subheader("Ask the agent")

question = st.text_input(
    "Question",
    placeholder=(
        "e.g. Explain the gradient, When is Test 1?, "
        "How do I find a tangent plane?"
    ),
)

example_col1, example_col2, example_col3, example_col4 = st.columns(4)

example_col1.caption("Try: Explain the gradient.")
example_col2.caption("Try: When is Test 1?")
example_col3.caption("Try: How do I find a tangent plane?")
example_col4.caption("Try: What happens if I miss a written test?")


if st.button(
    "Ask",
    type="primary",
    key="student_ask",
):
    if not question.strip():
        st.warning("Please enter a question.")

    else:
        with st.spinner("Searching the approved course materials..."):
            retrieved = retrieve(
                question,
                chunks,
            )

            answer = call_ai(
                question,
                retrieved,
            )

        if answer:
            st.markdown(answer)
        else:
            st.markdown(
                retrieval_answer(
                    question,
                    retrieved,
                )
            )

        if retrieved:
            with st.expander("Sources used"):
                for item in retrieved:
                    st.markdown(
                        f"**{item['source']} — {item['location']}** "
                        f"(relevance {item['score']:.2f})"
                    )

                    st.caption(
                        item["text"][:900]
                    )


# ============================================================
# PROGRAMME SCHEDULE
# ============================================================

st.divider()
st.subheader("Programme schedule")

if schedule is not None and not schedule.empty:
    display_schedule = schedule.copy()

    display_schedule["Date"] = display_schedule["Date"].apply(
        lambda d: d.strftime("%d %b %Y")
        if pd.notna(d)
        else ""
    )

    columns_to_show = [
        column
        for column in [
            "Date",
            "Item",
            "Details",
            "Trainer",
        ]
        if column in display_schedule.columns
    ]

    with st.expander(
        "View UPP Mathematics schedule",
        expanded=False,
    ):
        st.dataframe(
            display_schedule[columns_to_show],
            use_container_width=True,
            hide_index=True,
        )

    st.info(
        'You can ask questions such as "When is Test 1?", '
        '"When is the assignment due?", or '
        '"When is Partial Differentiation taught?"'
    )

else:
    st.info(
        "Programme schedule is currently unavailable."
    )


# ============================================================
# FOOTER
# ============================================================

st.divider()

st.caption(
    "Answers are based on the approved UPP Mathematics course materials "
    "included with this application."
)
